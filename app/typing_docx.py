import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.config import BASE_DIR

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'none'))
            if 'sz' in edge_data:
                element.set(qn('w:sz'), str(edge_data['sz']))
            if 'color' in edge_data:
                element.set(qn('w:color'), edge_data['color'])
            tcBorders.append(element)
        else:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), 'none')
            tcBorders.append(element)
    tcPr.append(tcBorders)


def build_typing_docx(filepath: str, text: str, language: str, set_num: int, date_str: str, word_count: int):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        header = section.header
        h_para = header.paragraphs[0]
        h_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = h_para.add_run(f"Learn with HiM • Typing with HiM • {date_str} • Set #{set_num:02d}")
        h_run.font.name = "Times New Roman"
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = RGBColor(148, 163, 184)

        footer = section.footer
        f_para = footer.paragraphs[0]
        f_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_para.add_run(
            "Insta: @Learnwithhimm  |  YT: @LearnwithHiM  |  TG: @Learnwithhim  |  TG Chat: @Learnwithhimm\n"
            "Official Paper-to-Screen Examination Material • Learn with HiM Assessment Cell"
        )
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(8.0)
        f_run.font.color.rgb = RGBColor(2, 132, 199)

    font_name = "Times New Roman" if language.lower() == "english" else "Mangal"
    logo_left_path = os.path.abspath(os.path.join(BASE_DIR, "assets", "logo.png"))
    logo_right_path = os.path.abspath(os.path.join(BASE_DIR, "assets", "logohim.png"))

    header_table = doc.add_table(rows=1, cols=3)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    widths = [Inches(1.0), Inches(4.8), Inches(1.0)]
    for row in header_table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(row.cells[idx])

    cell_left = header_table.cell(0, 0)
    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(logo_left_path):
        p_left.add_run().add_picture(logo_left_path, width=Inches(0.8))

    cell_mid = header_table.cell(0, 1)
    p_mid = cell_mid.paragraphs[0]
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_mid.add_run("Learn with HiM Typing Book\n")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    r_sub = p_mid.add_run("Type Daily! Type Smartly! Daily Free Relevant Typing Material!")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(22, 163, 74)

    cell_right = header_table.cell(0, 2)
    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(logo_right_path):
        p_right.add_run().add_picture(logo_right_path, width=Inches(0.8))

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_widths = [Inches(3.8), Inches(3.0)]

    for row in info_table.rows:
        for idx, width in enumerate(info_widths):
            cell = row.cells[idx]
            cell.width = width
            set_cell_border(
                cell,
                top={'sz': 4, 'val': 'single', 'color': 'CBD5E1'},
                bottom={'sz': 4, 'val': 'single', 'color': 'CBD5E1'},
                left={'sz': 4, 'val': 'single', 'color': 'CBD5E1'},
                right={'sz': 4, 'val': 'single', 'color': 'CBD5E1'}
            )

    cell_00 = info_table.cell(0, 0).paragraphs[0]
    r_00 = cell_00.add_run("Assessment: CAPF HCM / ASI STENO SKILL TEST")
    r_00.font.name = "Times New Roman"
    r_00.font.size = Pt(9)
    r_00.font.bold = True
    r_00.font.color.rgb = RGBColor(30, 58, 138)

    cell_01 = info_table.cell(0, 1).paragraphs[0]
    cell_01.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_01 = cell_01.add_run(f"Assessment Set: SET #{set_num:02d}")
    r_01.font.name = "Times New Roman"
    r_01.font.size = Pt(9)
    r_01.font.bold = True

    cell_10 = info_table.cell(1, 0).paragraphs[0]
    r_10 = cell_10.add_run(f"Language & Font: {language.upper()} ({font_name} 12pt)")
    r_10.font.name = "Times New Roman"
    r_10.font.size = Pt(8.5)
    r_10.font.color.rgb = RGBColor(71, 85, 105)

    cell_11 = info_table.cell(1, 1).paragraphs[0]
    cell_11.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_11 = cell_11.add_run(f"Date: {date_str}   |   Target Words: ~{word_count} Wds")
    r_11.font.name = "Times New Roman"
    r_11.font.size = Pt(8.5)
    r_11.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Paragraph rendering with explicit First-Line Tab Indent
    raw_paras = text.split("\n\n")
    for p_content in raw_paras:
        clean_p = p_content.replace("\t", "").strip()
        if not clean_p:
            continue
        p_body = doc.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body.paragraph_format.first_line_indent = Inches(0.5)  # Exam standard tab indent
        p_body.paragraph_format.line_spacing = 1.5
        p_body.paragraph_format.space_after = Pt(10)

        r_body = p_body.add_run(clean_p)
        r_body.font.name = font_name
        r_body.font.size = Pt(12)
        r_body.font.color.rgb = RGBColor(15, 23, 42)

    eval_table = doc.add_table(rows=1, cols=1)
    eval_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_eval = eval_table.cell(0, 0)
    c_eval.width = Inches(6.8)
    set_cell_border(c_eval, top={'sz': 6, 'val': 'dashed', 'color': '94A3B8'}, bottom={'sz': 6, 'val': 'dashed', 'color': '94A3B8'})

    p_eval = c_eval.paragraphs[0]
    r_eval = p_eval.add_run(
        "Candidate Name: _______________________    Roll No: __________________    "
        "Net WPM: ______    Accuracy: ______%"
    )
    r_eval.font.name = "Times New Roman"
    r_eval.font.size = Pt(9)
    r_eval.font.bold = True
    r_eval.font.color.rgb = RGBColor(51, 65, 85)

    doc.save(filepath)
    return filepath